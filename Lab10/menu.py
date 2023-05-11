import dearpygui.dearpygui as dpg
from dearpygui_ext import logger
from my_package.speed_calc import *
from my_package.price_calc import *
from docx import Document
dpg.create_context()
def my_func(sender, data):
    if(dpg.get_value('select') == 'GB'):
        total_memory = (dpg.get_value('Mem_amount') * 1024)
    else:
        total_memory = dpg.get_value('Mem_amount')
    # total_memory = dpg.get_value('Mem_amount') 

    #print(total_memory, "MB")
    dpg.set_value('hdd', round(hdd_speed(total_memory), 2))
    dpg.set_value('ssd', round(ssd_speed(total_memory), 2))
    dpg.set_value('flash', round(flash_speed(total_memory), 2))
    dpg.set_value('hddpr', hdd_price(total_memory))
    dpg.set_value('ssdpr', ssd_price(total_memory))
    dpg.set_value('flashpr', flash_price(total_memory))
# def saving_results(sender, data):
#    results = Document()
#    table = results.add_table(rows = 2, cols = 3)
#     hdd_cell = table.cell(0,0)
#     hdd_cell.text = 'HDD'
#    ssd_cell = table.cell(0,1)
#    ssd_cell.text = 'SSD'
#    flash_cell = table.cell(0,2)
#    flash_cell.text = 'Flash'
#    results.save()

    

with dpg.window(height = 600, width=800):

    dpg.add_input_int(label='Enter the amount', on_enter=True, tag='Mem_amount')
    dpg.add_combo(label='Select MB or GB', items=['MB', 'GB'], tag='select', callback=my_func)
    dpg.add_button(label = "Calculate ", callback = my_func)
    dpg.add_text('Reading/recording time in sec')
    dpg.add_input_text(label='HDD', callback=my_func, tag='hdd')
    dpg.add_input_text(label='SSD', callback=my_func, tag='ssd')
    dpg.add_input_text(label='Flash', callback=my_func, tag='flash')
    dpg.add_text('Prices 2nd task')
    dpg.add_input_text(label='HDD', callback=my_func, tag='hddpr')
    dpg.add_input_text(label='SSD', callback=my_func, tag='ssdpr')
    dpg.add_input_text(label='Flash', callback=my_func, tag='flashpr')

    # dpg.add_button(label='Save results', callback=saving_results)


    
dpg.create_viewport(title='Devices', width=800, height=600)
dpg.setup_dearpygui()
dpg.show_viewport()
dpg.start_dearpygui()
dpg.destroy_context()
